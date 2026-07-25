import glob, os
import pandas as pd
from openpyxl import load_workbook

files = sorted(glob.glob(r"d:\Work\Etiquette\Etiquette\data\**\*.xlsx", recursive=True))
for f in files:
    print("=" * 100)
    print(os.path.basename(f), f"({os.path.getsize(f)//1024} KB)")
    try:
        wb = load_workbook(f, read_only=True, data_only=True)
        for ws in wb.worksheets:
            print(f"  SHEET '{ws.title}': {ws.max_row} rows x {ws.max_column} cols")
        wb.close()
        xl = pd.ExcelFile(f)
        for sh in xl.sheet_names:
            df = xl.parse(sh, nrows=3)
            print(f"  --- {sh} columns ({len(df.columns)}):")
            print("     ", list(df.columns))
    except Exception as e:
        print("  ERR", type(e).__name__, e)
