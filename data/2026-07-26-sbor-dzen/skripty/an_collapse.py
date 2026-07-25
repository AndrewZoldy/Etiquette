"""Перепроверка обвала дистрибуции на данных выгрузки Студии."""
import os
import numpy as np
import pandas as pd
from openpyxl import load_workbook

pd.set_option("display.width", 220)
pd.set_option("display.max_columns", 40)

DATA = r"d:\Work\Etiquette\Etiquette\data"
STUDIO = os.path.join(DATA, "2026-07-25-dzen-vse-vremya.xlsx")
REN = {"Дата публикации": "date", "Час": "hour", "День недели": "dow", "Заголовок": "title",
       "Показы": "shows", "Открытия": "opens", "Дочитывания": "reads",
       "Время просмотра, мин": "minutes", "Лайки": "likes", "Комментарии": "comments",
       "Подписки": "subs", "CTR (открытия/показы)": "ctr", "Дочитываемость": "read_rate",
       "Минут на дочитывание": "min_per_read", "Лайки / дочитывания": "likes_per_read",
       "Комментарии / дочитывания": "comments_per_read", "Подписки / дочитывания": "subs_per_read"}


def load(sheet):
    wb = load_workbook(STUDIO, data_only=True)
    ws = wb[sheet]
    hdr = None
    for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if any(str(v) == "Дата публикации" for v in row if v is not None):
            hdr = i
            break
    headers = [c.value for c in ws[hdr]]
    lc = headers.index("Ссылка") + 1
    oids = {r: (ws.cell(row=r, column=lc).hyperlink.target or "").rstrip("/").split("/")[-1]
            for r in range(hdr + 1, ws.max_row + 1) if ws.cell(row=r, column=lc).hyperlink}
    wb.close()
    df = pd.read_excel(STUDIO, sheet_name=sheet, header=hdr - 1).dropna(how="all").reset_index(drop=True)
    df["_r"] = range(hdr + 1, hdr + 1 + len(df))
    df["oid"] = df["_r"].map(oids)
    df = df.rename(columns=REN)
    df["date"] = pd.to_datetime(df["date"])
    df["ym"] = df["date"].dt.to_period("M")
    df["age_days"] = (pd.Timestamp("2026-07-25") - df["date"]).dt.days
    return df


art = load("Статьи")
rol = load("Ролики")
pos = load("Посты")
print(f"статей {len(art)}, роликов {len(rol)}, постов {len(pos)}")
print(f"период: {art['date'].min().date()} — {art['date'].max().date()}\n")

# ---------- 1. Месячная динамика статей ----------
g = art.groupby("ym").agg(
    n=("shows", "size"),
    med_shows=("shows", "median"), sum_shows=("shows", "sum"),
    med_opens=("opens", "median"),
    med_reads=("reads", "median"), sum_reads=("reads", "sum"),
    med_subs=("subs", "median"), sum_subs=("subs", "sum"),
    med_ctr=("ctr", "median"), med_rr=("read_rate", "median"),
    med_lpr=("likes_per_read", "median"), med_cpr=("comments_per_read", "median"),
).round(4)
print("=" * 130)
print("СТАТЬИ по месяцу публикации")
print(g.to_string())

# ---------- 2. Ролики ----------
gr = rol.groupby("ym").agg(n=("shows", "size"), med_shows=("shows", "median"),
                           sum_shows=("shows", "sum"), med_opens=("opens", "median"),
                           med_ctr=("ctr", "median")).round(4)
print("\n" + "=" * 130)
print("РОЛИКИ по месяцу публикации")
print(gr.to_string())

# ---------- 3. Проверка поправки: сентябрь или октябрь ----------
print("\n" + "=" * 130)
print("ПРОВЕРКА ПОПРАВКИ (пик vs дно)")
peak = art[(art.ym >= "2025-01") & (art.ym <= "2025-07")]
bottom = art[(art.ym >= "2025-11") & (art.ym <= "2026-06")]
for name, d in (("пик 2025-01..07", peak), ("дно 2025-11..2026-06", bottom)):
    print(f"{name}: n={len(d)} медиана показов={d.shows.median():,.0f} "
          f"сумма показов={d.shows.sum():,.0f} медиана дочитываний={d.reads.median():,.0f} "
          f"сумма дочитываний={d.reads.sum():,.0f}")
print(f"\nмедианные показы упали в {peak.shows.median()/bottom.shows.median():.1f} раз")
print(f"суммарные показы(в месяц) упали в "
      f"{(peak.shows.sum()/7)/(bottom.shows.sum()/8):.1f} раз")
print(f"медианные дочитывания упали в {peak.reads.median()/bottom.reads.median():.1f} раз")

# ---------- 4. Концентрация ----------
print("\n" + "=" * 130)
print("КОНЦЕНТРАЦИЯ РАЗДАЧИ: доля статей с >=1 млн показов, по полугодиям")
art["half"] = art["date"].dt.year.astype(str) + "H" + ((art["date"].dt.month > 6).astype(int) + 1).astype(str)
c = art.groupby("half").apply(lambda d: pd.Series({
    "n": len(d),
    "доля>=1млн": round((d.shows >= 1e6).mean(), 3),
    "доля>=100тыс": round((d.shows >= 1e5).mean(), 3),
    "доля<10тыс": round((d.shows < 1e4).mean(), 3),
    "медиана": int(d.shows.median()),
    "сумма_млн": round(d.shows.sum() / 1e6, 1),
    "топ10%_доля_показов": round(d.shows.nlargest(max(1, len(d)//10)).sum()/d.shows.sum(), 3),
}), include_groups=False)
print(c.to_string())

# ---------- 5. Возраст не объясняет ----------
print("\n" + "=" * 130)
print("ВОЗРАСТ НЕ ОБЪЯСНЯЕТ: сравнение когорт по возрасту публикации")
for lo, hi, lbl in ((0, 60, "0-60 дней"), (150, 220, "150-220 дней"),
                    (300, 400, "300-400 дней"), (400, 550, "400-550 дней")):
    d = art[(art.age_days >= lo) & (art.age_days < hi)]
    if len(d) >= 8:
        print(f"{lbl:14s} n={len(d):3d}  публ.{d.date.min().date()}..{d.date.max().date()}  "
              f"медиана показов={d.shows.median():>10,.0f}  дочитываний={d.reads.median():>8,.0f}")

# ---------- 6. Объём публикаций ----------
print("\n" + "=" * 130)
print("ОБЪЁМ ПУБЛИКАЦИЙ в месяц (статьи / ролики / посты)")
vol = pd.DataFrame({"статьи": art.groupby("ym").size(), "ролики": rol.groupby("ym").size(),
                    "посты": pos.groupby("ym").size()}).fillna(0).astype(int)
vol["всего"] = vol.sum(axis=1)
print(vol.tail(30).to_string())

# ---------- 7. Доли через границу ----------
print("\n" + "=" * 130)
print("ДОЛИ (CTR, дочитываемость) через границу обвала — гладко или ломается?")
print(g[["n", "med_ctr", "med_rr", "med_lpr", "med_cpr"]].tail(24).to_string())

art.to_pickle("out/art.pkl"); rol.to_pickle("out/rol.pkl"); pos.to_pickle("out/pos.pkl")
print("\nсохранено out/art.pkl, out/rol.pkl, out/pos.pkl")
