# -*- coding: utf-8 -*-
import openpyxl, os, glob, json
paths = [r'D:/Work/Etiquette/Etiquette/data/2026-07-25-dzen-vse-vremya.xlsx',
         r'D:/Work/Etiquette/Etiquette/data/Дзен_студия_выгрузка_все_время_20260725.xlsx']
for p in paths:
    print('='*100); print('FILE', p, os.path.getsize(p))
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    print('sheets:', wb.sheetnames)
    for sn in wb.sheetnames:
        ws = wb[sn]
        print('-'*80)
        print('SHEET', repr(sn), 'dims', ws.max_row, 'x', ws.max_column)
        rows = []
        for i, r in enumerate(ws.iter_rows(values_only=True)):
            rows.append(r)
            if i >= 6: break
        for i, r in enumerate(rows):
            print('  r%d:' % i, [str(x)[:40] if x is not None else None for x in r])
    wb.close()
print()
print('=== drive-download dir ===')
for p in glob.glob(r'D:/Work/Etiquette/Etiquette/data/drive-download-20260725T221121Z-1-001/*'):
    print(os.path.getsize(p), p)
