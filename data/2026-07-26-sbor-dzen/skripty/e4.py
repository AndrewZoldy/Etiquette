# -*- coding: utf-8 -*-
import openpyxl, os, glob, warnings
warnings.filterwarnings('ignore')
p = r'D:/Work/Etiquette/Etiquette/data/2026-07-25-dzen-vse-vremya.xlsx'
wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
ws = wb['Как читать']
print('=== ПОЛНЫЙ ЛИСТ "Как читать" ===')
for r in ws.iter_rows(values_only=True):
    print(' | '.join('' if x is None else str(x) for x in r))
wb.close()
print()
print('='*100)
for p2 in glob.glob(r'D:/Work/Etiquette/Etiquette/data/drive-download-20260725T221121Z-1-001/*.xlsx'):
    wb = openpyxl.load_workbook(p2, read_only=True, data_only=True)
    print('FILE', os.path.basename(p2), 'sheets:', wb.sheetnames)
    for sn in wb.sheetnames:
        ws = wb[sn]
        print('   SHEET', repr(sn), ws.max_row, 'x', ws.max_column)
        for i, r in enumerate(ws.iter_rows(values_only=True)):
            if i > 4: break
            print('     ', [str(x)[:45] if x is not None else None for x in r][:25])
    wb.close()
    print('-'*90)
