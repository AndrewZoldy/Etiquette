import os
from openpyxl import load_workbook

DATA = r"d:\Work\Etiquette\Etiquette\data"
for fn in ("2026-07-25-dzen-vse-vremya.xlsx",
           os.path.join("drive-download-20260725T221121Z-1-001", "Дзен_студия_выгрузка_все_время_20260725.xlsx")):
    f = os.path.join(DATA, fn)
    if not os.path.exists(f):
        print("нет файла", f)
        continue
    print("=" * 80)
    print(fn)
    wb = load_workbook(f, data_only=True)
    for ws in wb.worksheets:
        n = 0
        samples = []
        for row in ws.iter_rows():
            for c in row:
                if c.hyperlink is not None:
                    n += 1
                    if len(samples) < 4:
                        samples.append((c.coordinate, c.value, c.hyperlink.target))
        print(f"  {ws.title}: hyperlinks={n}")
        for s in samples:
            print("     ", s)
        if ws.title == "Статистика":
            print("     first rows:")
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                print("      ", [str(x)[:40] for x in row][:12])
                if i > 6:
                    break
    wb.close()
