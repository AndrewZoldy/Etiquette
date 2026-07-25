"""Единая таблица: статистика Студии (по id из гиперссылок) + карточки + тексты + иллюстрации + комментарии."""
import glob, json, os, re, unicodedata
import pandas as pd
from openpyxl import load_workbook

DATA = r"d:\Work\Etiquette\Etiquette\data"
STUDIO = os.path.join(DATA, "2026-07-25-dzen-vse-vremya.xlsx")
OUT = "out"
os.makedirs(OUT, exist_ok=True)

COLS = ["Дата публикации", "Час", "День недели", "Заголовок", "Ссылка", "Показы", "Открытия",
        "Дочитывания", "Время просмотра, мин", "Лайки", "Комментарии", "Подписки",
        "CTR (открытия/показы)", "Дочитываемость", "Минут на дочитывание",
        "Лайки / дочитывания", "Комментарии / дочитывания", "Подписки / дочитывания"]

REN = {"Дата публикации": "date", "Час": "hour", "День недели": "dow", "Заголовок": "title_studio",
       "Показы": "shows", "Открытия": "opens", "Дочитывания": "reads",
       "Время просмотра, мин": "minutes", "Лайки": "likes", "Комментарии": "comments",
       "Подписки": "subs", "CTR (открытия/показы)": "ctr", "Дочитываемость": "read_rate",
       "Минут на дочитывание": "min_per_read", "Лайки / дочитывания": "likes_per_read",
       "Комментарии / дочитывания": "comments_per_read", "Подписки / дочитывания": "subs_per_read"}


def load_sheet(sheet):
    """Прочитать лист + вытащить oid из гиперссылок колонки «Ссылка»."""
    wb = load_workbook(STUDIO, data_only=True)
    ws = wb[sheet]
    # найти строку заголовков
    hdr_row = None
    for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        vals = [str(v) for v in row if v is not None]
        if any("Дата публикации" == v for v in vals):
            hdr_row = i
            break
    headers = [c.value for c in ws[hdr_row]]
    link_col = headers.index("Ссылка") + 1
    oids = {}
    for r in range(hdr_row + 1, ws.max_row + 1):
        c = ws.cell(row=r, column=link_col)
        if c.hyperlink is not None:
            oids[r] = (c.hyperlink.target or "").rstrip("/").split("/")[-1]
    wb.close()

    df = pd.read_excel(STUDIO, sheet_name=sheet, header=hdr_row - 1)
    df = df.dropna(how="all").reset_index(drop=True)
    df["_xlrow"] = range(hdr_row + 1, hdr_row + 1 + len(df))
    df["oid"] = df["_xlrow"].map(oids)
    df = df.rename(columns=REN)
    return df


arts = load_sheet("Статьи")
print("Студия статьи:", arts.shape, "с oid:", arts["oid"].notna().sum())

# ---------- публичные карточки ----------
pubs = pd.DataFrame([json.loads(l) for l in open(f"{OUT}/publications.jsonl", encoding="utf-8")])
pubs = pubs.rename(columns={"views": "views_public", "likes": "likes_public",
                            "comments": "comments_public", "title": "title_public"})
print("Публичные карточки:", pubs.shape)

# ---------- тексты ----------
rows = []
for f in glob.glob(f"{OUT}/pages/*.json"):
    d = json.load(open(f, encoding="utf-8"))
    imgs = d.get("images") or []
    rows.append({
        "oid": d["oid"], "n_words": d.get("n_words"), "n_chars": d.get("n_chars"),
        "n_paragraphs": d.get("n_paragraphs"), "n_images": d.get("n_images"),
        "n_headers": d.get("n_headers"), "n_links": len(d.get("links") or []),
        "n_quotes": len(d.get("quotes") or []), "n_list_items": len(d.get("list_items") or []),
        "publishTime": d.get("publishTime"), "modificationTime": d.get("modificationTime"),
        "has_caption": sum(1 for i in imgs if i.get("caption")),
    })
texts = pd.DataFrame(rows)
print("Тексты:", texts.shape)

m = arts.merge(pubs, on="oid", how="outer", indicator="src").merge(texts, on="oid", how="left")
print("\nсклейка:", m.shape)
print(m["src"].value_counts().to_dict())

# сколько статей Студии не нашли публичной карточки
missing_pub = m[m["src"] == "left_only"][["date", "title_studio", "shows", "reads"]]
print("\nв Студии есть, публичной карточки нет:", len(missing_pub))
print(missing_pub.head(10).to_string())
extra_pub = m[m["src"] == "right_only"][["title_public", "publication_date", "views_public"]]
print("\nкарточка есть, в Студии нет:", len(extra_pub))
print(extra_pub.head(10).to_string())

m.to_pickle(f"{OUT}/master_raw.pkl")
print("\nсохранено out/master_raw.pkl")
print("\nколонки:", [c for c in m.columns])
